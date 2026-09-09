"""Read-only vendor workbook evidence, keeping source cells and caches distinct."""
from datetime import datetime

import openpyxl
import pandas as pd


def audit_vendor_workbook(path):
    """Return nonempty source cells, native formats/comments and metric date coverage.

    No formula evaluation, source correction, or workbook save occurs. A formula with
    no cached result stays a formula; it is not classified as absent source input.
    """
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    rows = []
    try:
        for sheet in formulas:
            for line in sheet:
                date = sheet.cell(line[0].row, 1).value
                for cell in line:
                    if cell.value is None and cell.comment is None:
                        continue
                    value = cached[sheet.title][cell.coordinate].value
                    rows.append(dict(sheet=sheet.title, cell=cell.coordinate,
                        date=date if isinstance(date, datetime) else None,
                        column=cell.column_letter, source_value=cell.value, cached_value=value,
                        formula=cell.data_type == 'f', number_format=cell.number_format,
                        comment=cell.comment.text if cell.comment else None))
        cells = pd.DataFrame(rows)
        tokens = cells.source_value.astype(str).str.strip().str.lower().isin(['null', 'nan', 'nan%'])
        cells['missing_token'] = tokens
        cells['cache_absent'] = cells.formula & cells.cached_value.isna()
        values = cells[cells.date.notna() & ~cells.formula & ~tokens].copy()
        values['numeric'] = pd.to_numeric(values.source_value, errors='coerce')
        coverage = values[values.numeric.notna()].groupby(['sheet', 'column']).agg(
            first_date=('date', 'min'), last_date=('date', 'max'), numeric_dates=('date', 'nunique')).reset_index()
        return cells, coverage
    finally:
        formulas.close()
        cached.close()
