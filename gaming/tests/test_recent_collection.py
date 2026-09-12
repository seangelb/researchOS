"""Recent collection is tested only against retained fixtures and temporary storage."""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pandas as pd
import pytest
import requests
import sqlite3
from openpyxl import load_workbook

from variant_gaming import recent
from variant_gaming.common import save_raw_bytes, sha256_bytes
from variant_gaming.consolidate import build_state_period_revenue
from variant_gaming.storage import connect_readonly

FIXTURES = Path(__file__).parent / "fixtures"
NY_HTML = '''<h2>Sports Wagering</h2><table>
<tr><td>STATEWIDE</td><td><a href="/statewide.xlsx">Statewide Sports Wagering Weekly Excel</a></td></tr>
<tr><td>FanDuel</td><td><a href="/fanduel.xlsx">FanDuel Weekly Excel</a></td></tr></table>'''
MA_HTML = '''<h2>2023</h2><a href="/March-Rev-Report.pdf">March Revenue</a>
<h2>2026</h2><a href="/MGC-Revenue-Report-July-2026.pdf">July Revenue</a>'''


@pytest.fixture(autouse=True)
def no_live_http(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("Tests must not make live HTTP requests")
    monkeypatch.setattr(requests.sessions.Session, "request", denied)


def fake_session(state, content=None):
    if content is None:
        fixture = (FIXTURES / "NY/sample_statewide_weekly.xlsx" if state == "NY"
                   else FIXTURES / "MA/MGC-Revenue-Report-July-2026.pdf")
        content = fixture.read_bytes()
    sess = Mock()

    def get(url, **kwargs):
        if url in [recent.ma.LANDING_URL, recent.ma.ARCHIVE_URL, recent.ny.LANDING_URL]:
            return SimpleNamespace(text=NY_HTML if state == "NY" else MA_HTML, raise_for_status=lambda: None)
        return SimpleNamespace(content=content, url=url, raise_for_status=lambda: None)
    sess.get.side_effect = get
    return sess


def run(tmp_path, state, session, **kwargs):
    return recent.collect_recent(state, root=tmp_path, db_path=tmp_path / "candidate.sqlite",
                                 session=session, **kwargs)


def stored(tmp_path):
    conn = connect_readonly(tmp_path / "candidate.sqlite")
    try:
        return pd.read_sql_query("SELECT * FROM gaming_results", conn)
    finally:
        conn.close()


def test_ma_recent_discovery_fetches_only_latest_pdf(tmp_path):
    sess = fake_session("MA")
    log = run(tmp_path, "MA", sess, report_limit=1)
    assert log.report.tolist() == ["2026-07"]
    assert log.validation.tolist() == ["passed"]
    assert log.parsed_rows.tolist() == [8]
    assert log.period_start.tolist() == ["2026-07-01"]
    urls = [call.args[0] for call in sess.get.call_args_list]
    assert len(urls) == 3  # two index pages and one PDF; no historical report downloads
    assert not any("March-Rev" in url for url in urls)
    rows = stored(tmp_path)
    assert rows.loc[rows.operator == "STATEWIDE", "gross_revenue"].iloc[0] == 66530607.93


def test_ny_discovery_keeps_native_names_and_weekly_periods(tmp_path):
    log = run(tmp_path, "NY", fake_session("NY"))
    assert log.report.tolist() == ["STATEWIDE", "FanDuel"]
    assert (log.validation == "passed").all()
    assert (log.parsed_rows == 2).all()
    assert set(stored(tmp_path).frequency) == {"weekly"}
    assert (stored(tmp_path).gross_revenue < 0).any()


@pytest.mark.parametrize("state", ["MA", "NY"])
def test_unchanged_skips_only_retained_and_ingested(tmp_path, monkeypatch, state):
    sess = fake_session(state)
    opts = {"report_limit": 1, "include_operators": False}
    first = run(tmp_path, state, sess, **opts)
    before = stored(tmp_path)
    parse = Mock(side_effect=AssertionError("Unchanged report must skip parsing"))
    monkeypatch.setattr(recent, "_parse_report", parse)
    second = run(tmp_path, state, sess, **opts)
    assert second.validation.tolist() == ["previously_ingested"]
    assert second.bytes_changed.tolist() == [False]
    assert second.parsed_rows.tolist() == second.stored_rows.tolist() == [0]
    assert second.existing_rows.tolist() == first.stored_rows.tolist()
    assert not parse.called
    pd.testing.assert_frame_equal(before, stored(tmp_path))


def test_retained_bytes_without_observations_still_parse(tmp_path):
    content = (FIXTURES / "NY/sample_statewide_weekly.xlsx").read_bytes()
    retained = save_raw_bytes(tmp_path, "NY", content, "saved.xlsx")
    log = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert log.validation.tolist() == ["passed"]
    assert log.stored_rows.tolist() == [2]
    assert log.source_file.tolist() == [retained.relative_to(tmp_path).as_posix()]


def test_missing_retained_file_does_not_skip_parse(tmp_path, monkeypatch):
    first = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    (tmp_path / first.iloc[0].source_file).unlink()  # only this test's temporary capture
    parser = Mock(wraps=recent._parse_report)
    monkeypatch.setattr(recent, "_parse_report", parser)
    log = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert parser.call_count == 1
    assert log.validation.tolist() == ["passed"]


@pytest.mark.parametrize("state", ["MA", "NY"])
def test_force_reparse_preserves_first_capture(tmp_path, state):
    sess = fake_session(state)
    opts = {"include_operators": False, "report_limit": 1}
    first = run(tmp_path, state, sess, **opts)
    before = stored(tmp_path)
    log = run(tmp_path, state, sess, force_reparse=True, **opts)
    assert log.reason.tolist() == ["Forced reparse"]
    assert log.parsed_rows.tolist() == first.parsed_rows.tolist()
    pd.testing.assert_frame_equal(before, stored(tmp_path))


def revised_ny_bytes():
    book = load_workbook(FIXTURES / "NY/sample_statewide_weekly.xlsx")
    sheet = book.active
    header = next(row for row in sheet if any(cell.value == "Week-Ending" for cell in row))
    ggr_col = next(cell.column for cell in header if cell.value == "GGR")
    first_week = header[0].row + 1
    total = next(row[0].row for row in sheet if any(cell.value == "Total" for cell in row))
    sheet.cell(first_week, ggr_col).value += 100
    sheet.cell(total, ggr_col).value += 100
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def test_revised_workbook_retains_conflict_instead_of_latest_wins(tmp_path):
    run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    log = run(tmp_path, "NY", fake_session("NY", revised_ny_bytes()), include_operators=False)
    assert log.bytes_changed.tolist() == [True]
    assert log.validation.tolist() == ["passed"]
    rows = stored(tmp_path)
    assert len(rows) == 4 and rows.source_sha256.nunique() == 2
    assert rows.source_file.nunique() == 2
    periods = build_state_period_revenue(rows, metric="gross_revenue")
    conflict = periods[periods.completeness == "conflicting_sources"]
    assert len(conflict) == 1 and conflict.gross_revenue.isna().all()


def test_revised_ma_bytes_retained_separately(tmp_path):
    run(tmp_path, "MA", fake_session("MA"), report_limit=1)
    revised = (FIXTURES / "MA/MGC-Revenue-Report-July-2026.pdf").read_bytes() + b"\n"
    log = run(tmp_path, "MA", fake_session("MA", revised), report_limit=1)
    assert log.bytes_changed.tolist() == [True]
    assert stored(tmp_path).source_sha256.nunique() == 2


@pytest.mark.parametrize("failure", ["discovery", "download", "parse"])
def test_failures_are_visible_and_do_not_store_rows(tmp_path, failure):
    sess = fake_session("NY", b"not a workbook" if failure == "parse" else None)
    normal_get = sess.get.side_effect

    def get(url, **kwargs):
        if failure == "discovery" or (failure == "download" and url != recent.ny.LANDING_URL):
            raise requests.HTTPError("Fixture HTTP 503")
        return normal_get(url, **kwargs)
    sess.get.side_effect = get
    log = run(tmp_path, "NY", sess, include_operators=False)
    assert log.stored_rows.tolist() == [0]
    assert log.reason.str.len().gt(0).all()
    if failure == "discovery":
        assert log.download_status.tolist() == ["not_attempted"]
        assert not (tmp_path / "candidate.sqlite").exists()
    else:
        assert stored(tmp_path).empty
        assert log.download_status.tolist() == ["downloaded" if failure == "parse" else "failed"]
        if failure == "parse":
            assert log.validation.tolist() == ["failed"]
            assert (tmp_path / log.iloc[0].source_file).is_file()


def test_one_failure_does_not_hide_other_report_success(tmp_path):
    sess = fake_session("NY")
    original = sess.get.side_effect
    def get(url, **kwargs):
        if url.endswith("statewide.xlsx"):
            raise requests.HTTPError("Fixture failure")
        return original(url, **kwargs)
    sess.get.side_effect = get
    log = run(tmp_path, "NY", sess)
    assert log.download_status.tolist() == ["failed", "downloaded"]
    assert log.stored_rows.tolist() == [0, 2]


def test_unsupported_recent_and_bad_limit_have_no_side_effects(tmp_path):
    sess = fake_session("NY")
    for state, opts in [("NJ", {}), ("NY", {"vertical": "online_casino"}), ("MA", {"report_limit": 0})]:
        with pytest.raises(ValueError):
            run(tmp_path, state, sess, **opts)
    assert not sess.get.called
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("damage", ["duplicate", "date", "missing"])
def test_invalid_normalized_report_is_not_stored(tmp_path, monkeypatch, damage):
    parse = recent._parse_report
    def malformed(*args, **kwargs):
        frame = parse(*args, **kwargs)
        if damage == "duplicate":
            return pd.concat([frame, frame.iloc[[0]]], ignore_index=True)
        if damage == "date":
            frame.loc[0, "period_end"] = "2000-01-01"
        else:
            frame.loc[0, "gross_revenue"] = None
        return frame
    monkeypatch.setattr(recent, "_parse_report", malformed)
    log = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert log.validation.tolist() == ["failed"]
    assert stored(tmp_path).empty


@pytest.mark.parametrize("conflicting", [False, True])
def test_overlapping_ny_fiscal_sheets_cannot_hide_disagreement(tmp_path, conflicting):
    book = load_workbook(FIXTURES / "NY/sample_statewide_weekly.xlsx")
    copy = book.copy_worksheet(book.active)
    if conflicting:
        header = next(row for row in copy if any(cell.value == "Week-Ending" for cell in row))
        ggr_col = next(cell.column for cell in header if cell.value == "GGR")
        total = next(row[0].row for row in copy if any(cell.value == "Total" for cell in row))
        copy.cell(header[0].row + 1, ggr_col).value += 100
        copy.cell(total, ggr_col).value += 100
    stream = BytesIO()
    book.save(stream)
    log = run(tmp_path, "NY", fake_session("NY", stream.getvalue()), include_operators=False)
    if conflicting:
        assert log.validation.tolist() == ["failed"]
        assert "Conflicting duplicate weeks" in log.iloc[0].reason
        assert stored(tmp_path).empty
    else:
        assert log.stored_rows.tolist() == [2]


def test_partial_ingestion_is_repaired_before_unchanged_skip(tmp_path):
    run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    with sqlite3.connect(tmp_path / "candidate.sqlite") as connection:
        connection.execute("DELETE FROM gaming_results WHERE period_end='2025-07-06'")
    log = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert log.validation.tolist() == ["passed"]
    assert log.stored_rows.tolist() == [2]
    assert len(stored(tmp_path)) == 2


def test_legacy_ingestion_without_receipt_is_verified_once(tmp_path):
    run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    with sqlite3.connect(tmp_path / "candidate.sqlite") as connection:
        connection.execute("DELETE FROM recent_ingestions")
    verified = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert verified.validation.tolist() == ["passed"]
    skipped = run(tmp_path, "NY", fake_session("NY"), include_operators=False)
    assert skipped.validation.tolist() == ["previously_ingested"]
